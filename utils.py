"""Utilidades para generación de documentos (Word, PDF) y manejo de archivos."""

from __future__ import annotations

import io
import json
import os
import re
import zipfile
from html import escape
from datetime import datetime, timezone, timedelta
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

from config import EXPORTS_DIR


def sanitize_filename(name: str) -> str:
    cleaned = re.sub(r'[\\/*?:"<>|]', "", name)
    cleaned = cleaned.replace(" ", "_").strip()
    return cleaned or "documento"


def generar_nombre_archivo(estudiante: str, actividad_nombre: str) -> str:
    """Genera el nombre del archivo con formato Estudiante_retro_AI#"""
    codigo = get_activity_code(actividad_nombre)

    # Limpiamos el nombre del estudiante
    nombre_limpio = "".join(c for c in estudiante if c.isalnum() or c in " _-").strip().replace(" ", "_")
    
    return f"{nombre_limpio}_retro_{codigo}"


def _normalize_activity_name(name: str | None) -> str:
    normalized = (name or "").strip().lower()
    return re.sub(r"\s+", " ", normalized)


def get_activity_code(name: str | None) -> str:
    """Genera un código corto para el nombre del archivo (ej. AI1, PI, FI)."""
    lower_name = _normalize_activity_name(name)

    if not lower_name:
        return "Gen"
    if "proyecto integrador" in lower_name:
        return "PI"
    if re.search(r"\bforo de integraci[oó]n\b", lower_name):
        return "FI"
    if "actividad integradora" in lower_name:
        numeros = {"uno": 1, "dos": 2, "tres": 3, "cuatro": 4, "cinco": 5, "seis": 6}

        match = re.search(r'\d+', lower_name)
        if match:
            return f"AI{match.group()}"

        for palabra, num in numeros.items():
            if re.search(rf"\b{palabra}\b", lower_name):
                return f"AI{num}"

        return "AI"

    return "Gen"


def feedback_to_moodle_html(text: str, nombre_asesor: str = "", id_asesor: str = "") -> str:
    """Genera HTML con formato estricto y exacto para Moodle."""
    # Escudo preventivo: Si la IA junta el saludo con el texto en la misma línea, lo forzamos a separarse
    text = re.sub(r"^(Apreciable,\s*[a-zA-ZáéíóúÁÉÍÓÚñÑüÜ\s]+[.:;])\s+(.+)$", r"\1\n\n\2", text, flags=re.MULTILINE | re.IGNORECASE)
    
    lines = [line.strip() for line in text.split("\n") if line.strip()]
    html_lines: list[str] = []
    
    signature_lines = [
        "asesor virtual",
        "con afecto.",
        "cordialmente.",
        "atentamente.",
        "saludos cordiales."
    ]
    if nombre_asesor:
        signature_lines.append(nombre_asesor.strip().lower())
    if id_asesor:
        signature_lines.append(id_asesor.strip().lower())

    for i, line in enumerate(lines):
        # Limpiamos los hashes, pero MANTENEMOS los asteriscos vivos
        clean_line = line.replace("##", "").strip()
        # Creamos una versión en minúsculas sin asteriscos solo para validaciones de reglas lógicas
        lower_line = clean_line.lower().replace("**", "").replace("*", "").strip()
        es_grupo = re.match(r"^m\d{1,2}c\d{1,2}g\d{1,3}-\d{3}$", lower_line)
        
        if lower_line.startswith("apreciable") or lower_line.startswith("criterio ") or lower_line in signature_lines or es_grupo:
            # Quitamos asteriscos al momento de envolver en <strong> para no tener * sueltos
            texto_limpio = clean_line.replace("**", "").replace("*", "")
            html_lines.append(f"<p><strong>{escape(texto_limpio)}</strong></p>")
            if lower_line.startswith("apreciable") or lower_line in ["cordialmente.", "atentamente.", "con afecto."]:
                html_lines.append("<p> </p>")
        else:
            safe_line = escape(clean_line)
            # Detecta asteriscos SIMPLES y DOBLES y los convierte a negrita HTML
            safe_line = re.sub(r"\*{1,2}([^*]+)\*{1,2}", r'<strong style="font-size: 1rem;">\1</strong>', safe_line)
            # Detecta URLs y las convierte a enlaces
            safe_line = re.sub(r"(https?://[^\s]+)", r'<a href="\1">\1</a>', safe_line)
            
            html_lines.append(f'<p><span style="font-size: 1rem;">{safe_line}</span></p>')
            
            if i < len(lines) - 1:
                next_clean = lines[i+1].replace("##", "").strip().lower().replace("**", "").replace("*", "").strip()
                next_es_grupo = re.match(r"^m\d{1,2}c\d{1,2}g\d{1,3}-\d{3}$", next_clean)
                
                if not (lower_line in signature_lines and (next_clean in signature_lines or next_es_grupo)):
                    html_lines.append("<p> </p>")

    return "\n".join(html_lines)


def now_slug() -> str:
    tz_utc_minus_6 = timezone(timedelta(hours=-6))
    return datetime.now(tz_utc_minus_6).strftime("%Y%m%d_%H%M%S")


def set_run_font(run: Any, nombre: str = "Arial", tamano: int = 12, bold: bool = False, italic: bool = False, underline: bool = False) -> None:
    run.font.name = nombre
    run.font.size = Pt(tamano)
    run.font.bold = bold
    run.font.italic = italic
    run.font.underline = underline


def add_hyperlink(paragraph: Any, text: str, url: str) -> None:
    hyperlink = OxmlElement("w:hyperlink")
    r_id = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink.set(qn("r:id"), r_id)

    run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    rFonts = OxmlElement("w:rFonts")
    rFonts.set(qn("w:ascii"), "Arial")
    rFonts.set(qn("w:hAnsi"), "Arial")
    rPr.append(rFonts)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "24")
    rPr.append(sz)
    szCs = OxmlElement("w:szCs")
    szCs.set(qn("w:val"), "24")
    rPr.append(szCs)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0000FF")
    rPr.append(color)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rPr.append(u)
    run.append(rPr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def agregar_parrafo_firma(doc: Document, texto: str) -> Any:
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    
    run = p.add_run(texto)
    set_run_font(run, nombre="Arial", tamano=12, bold=True)
    return p


def add_formatted_line_to_doc(doc: Document, line: str, nombre_asesor: str = "", id_asesor: str = "") -> Any:
    stripped = line.strip()

    if not stripped:
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        return p

    signature_lines = [
        "asesor virtual",
        "con afecto.",
        "cordialmente.",
        "atentamente.",
        "saludos cordiales."
    ]
    if nombre_asesor:
        signature_lines.append(nombre_asesor.strip().lower())
    if id_asesor:
        signature_lines.append(id_asesor.strip().lower())
    
    lower_stripped = stripped.lower().replace("**", "").replace("*", "")
    es_grupo = re.match(r"^m\d{1,2}c\d{1,2}g\d{1,3}-\d{3}$", lower_stripped)

    if lower_stripped in signature_lines or es_grupo:
        texto_limpio = stripped.replace("**", "").replace("*", "")
        return agregar_parrafo_firma(doc, texto_limpio)

    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    known_headings = [
        "criterio cognitivo",
        "criterio actitudinal",
        "criterio comunicativo",
        "criterio colaborativo",
        "criterio pensamiento crítico",
        "retroalimentación formativa"
    ]

    if lower_stripped in known_headings or lower_stripped.startswith("criterio "):
        texto_limpio = stripped.replace("**", "").replace("*", "")
        run = p.add_run(texto_limpio)
        set_run_font(run, nombre="Arial", tamano=12, bold=True)
        return p

    if lower_stripped.startswith("apreciable"):
        texto_limpio = stripped.replace("**", "").replace("*", "")
        run = p.add_run(texto_limpio)
        set_run_font(run, nombre="Arial", tamano=12, bold=True)
        return p

    # Estandarizar asteriscos simples a dobles para que el generador de Word los procese igual
    normalized = re.sub(r"(?<!\*)\*([^*]+)\*(?!\*)", r"**\1**", stripped)
    normalized = normalized.replace("***", "**").replace("##", "").strip()
    
    tokens = re.split(r"(\*\*.*?\*\*|https?://[^\s]+)", normalized)

    for token in tokens:
        if not token:
            continue

        if token.startswith("**") and token.endswith("**") and len(token) >= 4:
            bold_text = token[2:-2]
            run = p.add_run(bold_text)
            set_run_font(run, nombre="Arial", tamano=12, bold=True)
        elif token.startswith("http://") or token.startswith("https://"):
            clean_url = token.rstrip(".,;)")
            suffix = token[len(clean_url):]
            add_hyperlink(p, clean_url, clean_url)
            if suffix:
                run = p.add_run(suffix)
                set_run_font(run, nombre="Arial", tamano=12, bold=False)
        else:
            clean_token = token.replace("**", "")
            run = p.add_run(clean_token)
            set_run_font(run, nombre="Arial", tamano=12, bold=False)

    return p


def docx_bytes(title: str, text: str, nombre_asesor: str = "", id_asesor: str = "") -> bytes:
    doc = Document()

    for section in doc.sections:
        section.top_margin = Pt(72)
        section.bottom_margin = Pt(72)
        section.left_margin = Pt(72)
        section.right_margin = Pt(72)

    # Forzar separación del saludo igual que en HTML por si la IA lo pegó
    text = re.sub(r"^(Apreciable,\s*[a-zA-ZáéíóúÁÉÍÓÚñÑüÜ\s]+[.:;])\s+(.+)$", r"\1\n\n\2", text, flags=re.MULTILINE | re.IGNORECASE)
    lines = text.split("\n")

    for line in lines:
        stripped = line.strip()

        if nombre_asesor and nombre_asesor in stripped and "Asesor virtual" in stripped:
            match_grupo = re.search(r"(M\d{1,2}C\d{1,2}G\d{1,3}-\d{3})", stripped, re.IGNORECASE)
            cohort = match_grupo.group(1).upper() if match_grupo else "M11C1G77-050"
            
            greeting = "Cordialmente." if "Cordialmente" in stripped else "Con afecto."
            partes_firma = [
                greeting,
                nombre_asesor,
                "Asesor virtual",
                id_asesor,
                cohort
            ]
            for parte in partes_firma:
                if parte:
                    agregar_parrafo_firma(doc, parte)
            continue

        add_formatted_line_to_doc(doc, line, nombre_asesor, id_asesor)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def pdf_bytes(title: str, text: str) -> bytes:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=54, leftMargin=54, topMargin=54, bottomMargin=54)
    styles = getSampleStyleSheet()
    normal_style = ParagraphStyle("CustomNormal", parent=styles["Normal"], fontName="Helvetica", fontSize=11, leading=16.5, spaceBefore=0, spaceAfter=0, alignment=4)
    title_style = ParagraphStyle("CustomTitle", parent=styles["Heading1"], fontName="Helvetica-Bold", fontSize=14, leading=18, spaceAfter=12, alignment=4)

    # Forzar separación del saludo
    text = re.sub(r"^(Apreciable,\s*[a-zA-ZáéíóúÁÉÍÓÚñÑüÜ\s]+[.:;])\s+(.+)$", r"\1\n\n\2", text, flags=re.MULTILINE | re.IGNORECASE)

    story = []
    if title:
        story.append(Paragraph(title, title_style))
        story.append(Spacer(1, 10))

    for paragraph in text.split("\n"):
        p_text = paragraph.strip().replace("##", "")
        if p_text:
            # Soporta tanto asterisco simple como doble para negritas
            p_text = re.sub(r'\*{1,2}([^*]+)\*{1,2}', r'<b>\1</b>', p_text)
            p_formatted = p_text.replace("\n", "<br/>")
            story.append(Paragraph(p_formatted, normal_style))

    doc.build(story)
    buffer.seek(0)
    return buffer.getvalue()


def create_zip(archivos: list[tuple[str, bytes]]) -> bytes:
    """Empaqueta una lista de archivos (nombre, contenido_en_bytes) en un archivo ZIP."""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, 'w', zipfile.ZIP_DEFLATED) as zf:
        for nombre_archivo, data in archivos:
            zf.writestr(nombre_archivo, data)
    buf.seek(0)
    return buf.getvalue()


def export_json(data: dict[str, Any], filename_prefix: str = "export") -> Path:
    EXPORTS_DIR.mkdir(parents=True, exist_ok=True)
    filepath = EXPORTS_DIR / f"{filename_prefix}_{now_slug()}.json"
    with open(filepath, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
    return filepath
